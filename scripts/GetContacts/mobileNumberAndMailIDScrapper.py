import re
import docx
import PyPDF2
import xlsxwriter

#Getting file location from user
location = str(input())

#Make decision based on extension of file
extension = location[-3:]
if extension=="ocx" :
    contentFile = docx.Document(location)
    chunk = ""
    for p in contentFile.paragraphs:
        chunk+=p.text    
elif extension=="txt":
    contentFile = open(location)
    chunk = contentFile.read()
    contentFile.close()
elif extension=="pdf":
    contentFile = open(location,'rb')
    reader = PyPDF2.PdfFileReader(contentFile)
    chunk=""
    for page in range(reader.numPages):
        chunk+=reader.getPage(page).extractText()    

#Regex objects for phone numbers
phoneRegex = re.compile(r'9\d{9}')
phoneRegex1 = re.compile(r'8\d{9}')
phoneRegex2 = re.compile(r'7\d{9}')
phoneRegex3 = re.compile(r'6\d{9}')

#Regex for email addressess
emailRegex = re.compile(r"[a-z0-9\.\-+_]+@[a-z0-9\.\-+_]+\.[a-z]+")

#Extracted phone and mail addresses

extractedNumbers = phoneRegex.findall(chunk)
extractedNumbers1 = phoneRegex1.findall(chunk)
extractedNumbers2 = phoneRegex2.findall(chunk)
extractedNumbers3 = phoneRegex3.findall(chunk)
extractedEmail = emailRegex.findall(chunk)

mail_list = extractedEmail
phoneNumbers_list = extractedNumbers+extractedNumbers1+extractedNumbers2+extractedNumbers3

#Preparing a excel sheet
workbook = xlsxwriter.Workbook('ExtractedList.xlsx') 
worksheet = workbook.add_worksheet() 
row = 0
col = 0

#Adding email in sheet
for ids in mail_list:
    worksheet.write(row,col,ids)
    row+=1

col+=1
row = 0

#Adding numbers to sheet
for numbers in phoneNumbers_list:
    worksheet.write(row,col,numbers)
    row+=1

workbook.close()    
    
